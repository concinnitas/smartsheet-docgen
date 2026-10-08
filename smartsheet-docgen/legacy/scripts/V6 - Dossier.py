import os
import smartsheet
from docx import Document
from tkinter import ttk, filedialog, messagebox
import tkinter as tk
from ttkthemes import ThemedTk
import json

# Default paths for configuration and mappings
script_dir = os.path.abspath(os.path.dirname(__file__))
config_file = os.path.join(script_dir, "config.json")
mappings_file = os.path.join(script_dir, "mappings.json")

# Load the saved API key from config file
def load_config():
    try:
        if os.path.isfile(config_file):
            with open(config_file, 'r') as file:
                return json.load(file)
    except json.JSONDecodeError as e:
        print(f"Error loading config: {e}")
    return {}

# Save the API key to the config file
def save_config(api_key):
    try:
        with open(config_file, 'w') as file:
            json.dump({"api_key": api_key}, file, indent=4)
    except Exception as e:
        print(f"Error saving config: {e}")

# Load JSON mappings (static only)
def load_mappings():
    if not os.path.isfile(mappings_file):
        return None
    try:
        with open(mappings_file, 'r') as file:
            return json.load(file)
    except json.JSONDecodeError as e:
        messagebox.showerror("Error", f"Error parsing mapping file: {e}")
    return None

# Save a new JSON mappings file with static_cell_mappings placeholder
def save_new_mappings(content=None):
    if not content:
        content = {"static_cell_mappings": {}, "fallback": ""}
    with open(mappings_file, 'w') as file:
        json.dump(content, file, indent=4)

# Replace static placeholders in the Word document
def replace_static_placeholders(doc, sheet_data, static_mappings):
    for placeholder, details in static_mappings.items():
        row_index, col_index = details.get("location", [0, 0])
        fallback = details.get("fallback", "")
        try:
            cell_value = sheet_data.rows[row_index].cells[col_index].value
            value = cell_value if cell_value is not None else fallback
            for paragraph in doc.paragraphs:
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, str(value))
            for table in doc.tables:
                for row in table.rows:
                    for cell_obj in row.cells:
                        if placeholder in cell_obj.text:
                            cell_obj.text = cell_obj.text.replace(placeholder, str(value))
        except IndexError:
            print(f"Invalid index for placeholder {placeholder}: {row_index}, {col_index}")
            continue

# Main processing function
def process_smartsheet_and_generate_doc(api_token, sheet_name, template_file, output_file, mappings):
    try:
        ss_client = smartsheet.Smartsheet(api_token)
        sheets = ss_client.Sheets.list_sheets(include_all=True).data
        sheet = next((s for s in sheets if s.name == sheet_name), None)
        if sheet is None:
            raise ValueError(f"Sheet '{sheet_name}' not found.")
        sheet_data = ss_client.Sheets.get_sheet(sheet.id)

        doc = Document(template_file)
        static_mappings = mappings.get("static_cell_mappings", {})
        replace_static_placeholders(doc, sheet_data, static_mappings)

        doc.save(output_file)
        messagebox.showinfo("Success", f"Document generated: {output_file}")
    except Exception as e:
        messagebox.showerror("Error", str(e))

# GUI Creation
def create_ui():
    config = load_config()
    root = ThemedTk(theme="plastik")
    root.title("Smartsheet Report Generator")
    root.geometry("600x400")

    api_key_var = tk.StringVar(value=config.get("api_key", ""))
    sheet_name_var = tk.StringVar()
    template_file_var = tk.StringVar()
    output_dir_var = tk.StringVar()

    mappings = load_mappings()
    if not mappings:
        if messagebox.askyesno("Mappings Missing", "Mappings file not found. Create a new one?"):
            save_new_mappings()
            mappings = load_mappings()

    ttk.Label(root, text="API Key").grid(row=0, column=0, padx=10, pady=10)
    ttk.Entry(root, textvariable=api_key_var, width=40).grid(row=0, column=1, pady=10)
    ttk.Button(root, text="Import Key", command=lambda: import_api_key()).grid(row=0, column=2)

    def import_api_key():
        path = filedialog.askopenfilename(filetypes=[("Text Files", "*.txt")])
        if path:
            with open(path, 'r') as f:
                api_key_var.set(f.read().strip())
            save_config(api_key_var.get())
            fetch_smartsheets()

    ttk.Label(root, text="Smartsheet Name").grid(row=1, column=0)
    smartsheet_dropdown = ttk.Combobox(root, textvariable=sheet_name_var, width=38)
    smartsheet_dropdown.grid(row=1, column=1)
    ttk.Button(root, text="Fetch Sheets", command=lambda: fetch_smartsheets()).grid(row=1, column=2)

    ttk.Label(root, text="Template File").grid(row=2, column=0)
    ttk.Entry(root, textvariable=template_file_var, width=40).grid(row=2, column=1)
    ttk.Button(root, text="Browse", command=lambda: template_file_var.set(filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")]))).grid(row=2, column=2)

    ttk.Label(root, text="Output Directory").grid(row=3, column=0)
    ttk.Entry(root, textvariable=output_dir_var, width=40).grid(row=3, column=1)
    ttk.Button(root, text="Browse", command=lambda: output_dir_var.set(filedialog.askdirectory())).grid(row=3, column=2)

    def execute():
        if not all([api_key_var.get(), sheet_name_var.get(), template_file_var.get(), output_dir_var.get()]):
            messagebox.showerror("Error", "All fields are required")
            return
        output_file = os.path.join(output_dir_var.get(), f"{sheet_name_var.get()}_report.docx")
        save_config(api_key_var.get())
        process_smartsheet_and_generate_doc(
            api_key_var.get(),
            sheet_name_var.get(),
            template_file_var.get(),
            output_file,
            mappings
        )

    ttk.Button(root, text="Generate Report", command=execute).grid(row=4, column=1, pady=20)

    def fetch_smartsheets():
        try:
            client = smartsheet.Smartsheet(api_key_var.get())
            sheets = client.Sheets.list_sheets(include_all=True).data
            smartsheet_dropdown['values'] = [s.name for s in sheets]
            if sheets:
                smartsheet_dropdown.current(0)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to fetch sheets: {e}")

    if api_key_var.get():
        fetch_smartsheets()

    root.mainloop()

if __name__ == "__main__":
    create_ui()