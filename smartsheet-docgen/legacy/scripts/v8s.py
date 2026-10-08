import os
import smartsheet
from docx import Document
from tkinter import ttk, filedialog, messagebox
import tkinter as tk
from ttkthemes import ThemedTk
import json
import time

# Default paths for configuration and mappings
script_dir = os.path.abspath(os.path.dirname(__file__))
config_file = os.path.join(script_dir, "config.json")
mappings_file = os.path.join(script_dir, "mappings.json")

# Load the saved API key from config file
def load_config():
    try:
        if os.path.isfile(config_file):
            with open(config_file, 'r') as f:
                return json.load(f)
    except json.JSONDecodeError as e:
        print(f"Error loading config: {e}")
    return {}

# Save the API key to the config file
def save_config(api_key):
    try:
        with open(config_file, 'w') as f:
            json.dump({"api_key": api_key}, f, indent=4)
    except Exception as e:
        print(f"Error saving config: {e}")

# Load JSON mappings (static only)
def load_mappings():
    if not os.path.isfile(mappings_file):
        return None
    try:
        with open(mappings_file, 'r') as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        messagebox.showerror("Error", f"Error parsing mapping file: {e}")
    return None

# Save a new JSON mappings file with static_cell_mappings placeholder
def save_new_mappings(content=None):
    if not content:
        content = {"static_cell_mappings": {}}
    with open(mappings_file, 'w') as f:
        json.dump(content, f, indent=4)

# Replace static placeholders in the Word document
def replace_static_placeholders(doc, sheet_data, static_mappings):
    col_title_to_index = {col.title: idx for idx, col in enumerate(sheet_data.columns)}
    for placeholder, details in static_mappings.items():
        loc = details.get("location", {})
        fallback = details.get("fallback", "")
        try:
            if isinstance(loc, dict):
                row_idx = loc.get("row", 1) - 1
                col_name = loc.get("column")
                col_idx = col_title_to_index.get(col_name)
                if col_idx is None:
                    print(f"Column '{col_name}' not found for placeholder {placeholder}")
                    continue
            else:
                row_idx, col_idx = loc
            cell_value = sheet_data.rows[row_idx].cells[col_idx].value
            value = cell_value if cell_value is not None else fallback
            for paragraph in doc.paragraphs:
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, str(value))
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if placeholder in cell.text:
                            cell.text = cell.text.replace(placeholder, str(value))
        except Exception as e:
            print(f"Error replacing placeholder {placeholder}: {e}")

# Main processing function
def process_smartsheet_and_generate_doc(api_token, sheet_name, template_file, output_file, mappings):
    try:
        ss_client = smartsheet.Smartsheet(api_token)
        all_sheets = ss_client.Sheets.list_sheets(include_all=True).data
        sheet = next((s for s in all_sheets if s.name == sheet_name), None)
        if sheet is None:
            raise ValueError(f"Sheet '{sheet_name}' not found.")
        sheet_data = ss_client.Sheets.get_sheet(sheet.id)
        doc = Document(template_file)
        static_mappings = mappings.get("static_cell_mappings", {})
        replace_static_placeholders(doc, sheet_data, static_mappings)
        doc.save(output_file)
        messagebox.showinfo("Success", f"Document generated: {output_file}")
    except Exception as e:
        messagebox.showerror("Error", str(e))

# GUI Creation
def create_ui():
    config = load_config()
    root = ThemedTk(theme="plastik")
    root.title("Smartsheet Report Generator")
    root.geometry("620x560")

    # Variables
    api_key_var = tk.StringVar(value=config.get("api_key", ""))
    sheet_name_var = tk.StringVar()
    name_var = tk.StringVar()
    id_var = tk.StringVar()
    template_file_var = tk.StringVar()
    output_dir_var = tk.StringVar()
    output_filename_var = tk.StringVar()

    # State
    smartsheet_dropdown_entries = {}
    selected_row_idx = {'idx': None}

    # Load mappings
    mappings = load_mappings()
    if not mappings:
        if messagebox.askyesno("Mappings Missing", "Mappings file not found. Create a new one?" ):
            save_new_mappings()
            mappings = load_mappings()

    # Helper functions
    def fetch_smartsheets():
        try:
            client = smartsheet.Smartsheet(api_key_var.get())
            sheets = client.Sheets.list_sheets(include_all=True).data
            names = [s.name for s in sheets]
            smartsheet_dropdown.config(values=names)
            if names:
                smartsheet_dropdown.current(0)
                load_row_selectors()
        except Exception as e:
            messagebox.showerror("Error", f"Failed to fetch sheets: {e}")

    def import_api_key():
        path = filedialog.askopenfilename(filetypes=[("Text Files", "*.txt")])
        if path:
            with open(path, 'r') as f:
                api_key_var.set(f.read().strip())
            save_config(api_key_var.get())
            fetch_smartsheets()

    def load_row_selectors():
        try:
            client = smartsheet.Smartsheet(api_key_var.get())
            all_sheets = client.Sheets.list_sheets(include_all=True).data
            sheet = next((s for s in all_sheets if s.name == sheet_name_var.get()), None)
            if not sheet:
                return
            sheet_data = client.Sheets.get_sheet(sheet.id)
            col_map = {col.title: idx for idx, col in enumerate(sheet_data.columns)}
            if "Name" not in col_map or "ID Number" not in col_map:
                messagebox.showerror("Error", "Required columns 'Name' and/or 'ID Number' not found in sheet.")
                name_dropdown.config(values=[], state="disabled")
                id_dropdown.config(values=[], state="disabled")
                return
            name_idx = col_map["Name"]
            uid_idx = col_map["ID Number"]
            smartsheet_dropdown_entries.clear()
            for idx, row in enumerate(sheet_data.rows):
                name = row.cells[name_idx].value or ""
                uid = row.cells[uid_idx].value or ""
                smartsheet_dropdown_entries.setdefault(name, []).append((uid, idx))
            names = sorted(smartsheet_dropdown_entries.keys())
            name_dropdown.config(values=names, state="readonly")
            name_dropdown.set("")
            id_dropdown.set("")
            id_dropdown.config(values=[], state="disabled")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load rows: {e}")

    def on_name_selected(event=None):
        entries = smartsheet_dropdown_entries.get(name_var.get(), [])
        uids = [uid for uid, _ in entries]
        id_dropdown.config(values=uids, state="readonly" if uids else "disabled")
        if uids:
            id_dropdown.current(0)
            on_id_selected()

    def on_id_selected(event=None):
        entries = smartsheet_dropdown_entries.get(name_var.get(), [])
        for uid, idx in entries:
            if uid == id_var.get():
                selected_row_idx['idx'] = idx
                break

    def execute():
        if not all([
            api_key_var.get(),
            sheet_name_var.get(),
            template_file_var.get(),
            output_dir_var.get(),
            output_filename_var.get(),
            selected_row_idx['idx'] is not None
        ]):
            messagebox.showerror("Error", "All fields including row selection are required")
            return
        filename = output_filename_var.get()
        if not filename.lower().endswith(".docx"):
            filename += ".docx"
        output_file = os.path.join(output_dir_var.get(), filename)
        for details in mappings.get("static_cell_mappings", {}).values():
            loc = details.get("location")
            if isinstance(loc, dict):
                loc["row"] = selected_row_idx['idx'] + 1
        process_smartsheet_and_generate_doc(
            api_key_var.get(),
            sheet_name_var.get(),
            template_file_var.get(),
            output_file,
            mappings
        )

    # UI Elements
    ttk.Label(root, text="API Key").grid(row=0, column=0, padx=10, pady=10, sticky="e")
    ttk.Entry(root, textvariable=api_key_var, width=40).grid(row=0, column=1)
    ttk.Button(root, text="Import Key", command=import_api_key).grid(row=0, column=2)

    ttk.Label(root, text="Smartsheet Name").grid(row=1, column=0, sticky="e")
    smartsheet_dropdown = ttk.Combobox(root, textvariable=sheet_name_var, width=38)
    smartsheet_dropdown.grid(row=1, column=1)
    ttk.Button(root, text="Fetch Sheets", command=fetch_smartsheets).grid(row=1, column=2)
    smartsheet_dropdown.search_buffer = ""
    smartsheet_dropdown.last_search_time = time.time()
    smartsheet_dropdown.bind('<KeyRelease>', lambda e: on_sheet_key(e))
    smartsheet_dropdown.bind('<<ComboboxSelected>>', lambda e: load_row_selectors())

    def on_sheet_key(event):
        char = event.char
        now = time.time()
        if now - smartsheet_dropdown.last_search_time > 1.0:
            smartsheet_dropdown.search_buffer = ""
        smartsheet_dropdown.last_search_time = now
        if event.keysym == 'BackSpace':
            smartsheet_dropdown.search_buffer = smartsheet_dropdown.search_buffer[:-1]
        elif char.isprintable() and not char.isspace():
            smartsheet_dropdown.search_buffer += char
        else:
            return
        for idx, val in enumerate(smartsheet_dropdown['values']):
            if val.lower().startswith(smartsheet_dropdown.search_buffer.lower()):
                smartsheet_dropdown.current(idx)
                load_row_selectors()
                break

    ttk.Label(root, text="Name").grid(row=2, column=0, sticky="e")
    name_dropdown = ttk.Combobox(root, textvariable=name_var, width=38, state="disabled")
    name_dropdown.grid(row=2, column=1)
    name_dropdown.bind('<<ComboboxSelected>>', on_name_selected)

    ttk.Label(root, text="ID Number").grid(row=3, column=0, sticky="e")
    id_dropdown = ttk.Combobox(root, textvariable=id_var, width=38, state="disabled")
    id_dropdown.grid(row=3, column=1)
    id_dropdown.bind('<<ComboboxSelected>>', on_id_selected)

    # Template file
    ttk.Label(root, text="Template File").grid(row=4, column=0, sticky="e")
    ttk.Entry(root, textvariable=template_file_var, width=40).grid(row=4, column=1)
    ttk.Button(
        root,
        text="Browse",
        command=lambda: template_file_var.set(
            filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
        )
    ).grid(row=4, column=2)

    # Output directory
    ttk.Label(root, text="Output Directory").grid(row=5, column=0, sticky="e")
    ttk.Entry(root, textvariable=output_dir_var, width=40).grid(row=5, column=1)
    ttk.Button(
        root,
        text="Browse",
        command=lambda: output_dir_var.set(
            filedialog.askdirectory()
        )
    ).grid(row=5, column=2)

    # Output filename
    ttk.Label(root, text="Output Filename").grid(row=6, column=0, sticky="e")
    ttk.Entry(root, textvariable=output_filename_var, width=40).grid(row=6, column=1)

    # Generate button
    ttk.Button(root, text="Generate Report", command=execute).grid(row=7, column=1, pady=20)

    root.mainloop()

if __name__ == "__main__":
    create_ui()
